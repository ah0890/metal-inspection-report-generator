"""Covering letter Word document (python-docx), mirroring the PDF layout."""

from __future__ import annotations

import base64
import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from reports.letter_content import LetterContent

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x5B, 0x67, 0x75)


def _para(container, text: str = "", *, bold=False, size=10.5, color=None,
          align=WD_ALIGN_PARAGRAPH.LEFT, after=0, before=0):
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return p


def _bottom_border(paragraph, color="1F3A5F", size=12) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def _build(c: LetterContent) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(22)
    sec.top_margin, sec.bottom_margin = Mm(14), Mm(18)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(0)

    # Centred letterhead: name, tagline, address/contact; optional logo on the left
    # (an empty right-hand column of equal width keeps the text centred).
    centre = WD_ALIGN_PARAGRAPH.CENTER
    has_logo = bool(c.company.logo_base64)
    if has_logo:
        head = doc.add_table(rows=1, cols=3)
        head.alignment = WD_TABLE_ALIGNMENT.CENTER
        head.autofit = False
        for cell, w in zip(head.rows[0].cells, (Mm(22), Mm(122), Mm(22))):
            cell.width = w
        try:
            run = head.cell(0, 0).paragraphs[0].add_run()
            run.add_picture(io.BytesIO(base64.b64decode(c.company.logo_base64)), height=Mm(15))
        except Exception:  # unreadable image: continue without it
            pass
        block = head.cell(0, 1)
        first = block.paragraphs[0]
    else:
        block = doc
        first = _para(doc)
    first.alignment = centre
    r = first.add_run(c.company.name)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(15 if has_logo else 17), NAVY
    if c.company.tagline.strip():
        _para(block, c.company.tagline, size=8.5, color=MUTED, align=centre)
    address = ", ".join(line.strip() for line in c.company.address.splitlines() if line.strip())
    for line in (address, c.company.contact_line):
        if line:
            _para(block, line, size=7.8, color=MUTED, align=centre)
    rule = _para(doc, after=14)
    _bottom_border(rule)

    # Reference and date on one line, date right-aligned via a tab stop.
    p = _para(doc, after=12)
    p.paragraph_format.tab_stops.add_tab_stop(Mm(166), alignment=2)  # 2 = right
    r = p.add_run("Ref. No.: ")
    r.bold = True
    p.add_run(c.letter_number + "\t")
    r = p.add_run("Date: ")
    r.bold = True
    p.add_run(c.date_text)

    if c.recipient_lines:
        _para(doc, "To,")
        for line in c.recipient_lines:
            _para(doc, line)
        _para(doc, after=6)

    _para(doc, f"Subject: {c.subject}", bold=True)
    for ref in c.reference_lines:
        _para(doc, ref, size=9.5, color=MUTED)
    _para(doc, c.salutation, before=10, after=6)
    for text in c.paragraphs:
        _para(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=7)

    _para(doc, c.closing, before=4)
    if c.for_company:
        _para(doc, c.for_company, bold=True)
    _para(doc, after=40)  # signature space
    if c.signatory_name:
        _para(doc, c.signatory_name, bold=True)
    _para(doc, c.signatory_designation)

    if c.enclosures:
        _para(doc, "Enclosures:", bold=True, size=9.5, before=14)
        for i, e in enumerate(c.enclosures, start=1):
            _para(doc, f"{i}.  {e}", size=9.5)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("  |  ".join(p for p in (c.company.name, c.company.contact_line) if p))
    run.font.size, run.font.color.rgb = Pt(7.5), MUTED
    return doc


def generate_letter_docx(content: LetterContent) -> bytes:
    buf = io.BytesIO()
    _build(content).save(buf)
    return buf.getvalue()
