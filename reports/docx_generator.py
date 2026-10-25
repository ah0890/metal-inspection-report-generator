"""Word (DOCX) version of the inspection report.

Mirrors the PDF layout using the same ReportContent, with compact A4 margins
so the document normally fits on one page when opened in Word.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from config import settings
from models.evaluation import EvaluationResult
from models.inspection import InspectionReport
from reports.content import ReportContent, build_content

NAVY = "1F3A5F"
LABEL_BG = "EDF0F4"
MUTED = RGBColor(0x5B, 0x67, 0x75)
STATUS_RGB = {
    "PASS": ("17663A", "E3F2E8"),
    "FAIL": ("A61B1B", "FBE5E3"),
    "REVIEW": ("9A4A00", "FDF0DC"),
    "REVIEW REQUIRED": ("9A4A00", "FDF0DC"),
    "NOT EVALUATED": ("5B6775", "FFFFFF"),
}
CONTENT_W_MM = 182


def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _write(cell, text: str, *, bold=False, size=8.5, color: RGBColor | None = None,
           align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.space_before = Pt(0)
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _table(doc, rows: int, widths_mm: list[float]):
    t = doc.add_table(rows=rows, cols=len(widths_mm))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row in t.rows:
        for cell, w in zip(row.cells, widths_mm):
            cell.width = Mm(w)
    return t


def _section(doc, title: str) -> None:
    t = _table(doc, 1, [CONTENT_W_MM])
    cell = t.cell(0, 0)
    _shade(cell, NAVY)
    _write(cell, title.upper(), bold=True, size=8.5, color=RGBColor(0xFF, 0xFF, 0xFF))
    _gap(doc, 1)


def _gap(doc, points: float = 5) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(points)


def _label_value_table(doc, rows, widths_mm) -> None:
    t = _table(doc, len(rows), widths_mm)
    for r, values in enumerate(rows):
        for col, text in enumerate(values):
            cell = t.cell(r, col)
            if col % 2 == 0:
                _shade(cell, LABEL_BG)
                _write(cell, text, bold=True, size=7.5, color=MUTED)
            else:
                _write(cell, text)


def _results_table(doc, items, details: list[str]) -> None:
    widths = [CONTENT_W_MM * f for f in (0.30, 0.24, 0.30, 0.16)]
    t = _table(doc, 1 + len(items) + len(details), widths)
    for col, h in enumerate(("TEST", "RESULT", "REQUIREMENT", "STATUS")):
        _shade(t.cell(0, col), LABEL_BG)
        _write(t.cell(0, col), h, bold=True, size=7.5, color=MUTED)
    for r, item in enumerate(items, start=1):
        _write(t.cell(r, 0), item.name + (" *" if item.mandatory else ""), bold=True)
        _write(t.cell(r, 1), item.result_text)
        req = item.requirement_text + (f"\n{item.note}" if item.note and item.status.value == "REVIEW" else "")
        _write(t.cell(r, 2), req)
        fg, bg = STATUS_RGB[item.status.value]
        _shade(t.cell(r, 3), bg)
        _write(t.cell(r, 3), item.status.value, bold=True, size=8, color=RGBColor.from_string(fg),
               align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, text in enumerate(details):
        r = 1 + len(items) + i
        merged = t.cell(r, 0).merge(t.cell(r, 3))
        _write(merged, text, size=7, color=MUTED)


def _build(c: ReportContent) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(14)
    sec.top_margin, sec.bottom_margin = Mm(12), Mm(12)
    sec.footer_distance = Mm(6)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8.5)
    normal.paragraph_format.space_after = Pt(0)

    # Header block
    head = _table(doc, 1, [120, 62])
    head.style = doc.styles["Normal Table"]  # header block without borders
    left = head.cell(0, 0)
    _write(left, settings.LAB_NAME.upper(), bold=True, size=7.5, color=MUTED)
    p = left.add_paragraph()
    r = p.add_run(c.title)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(16), RGBColor.from_string(NAVY)
    p = left.add_paragraph()
    r = p.add_run(settings.LAB_SUBTITLE)
    r.font.size, r.font.color.rgb = Pt(8), MUTED
    right = head.cell(0, 1)
    _write(right, f"Report No.: {c.report_number}", bold=True, size=8.5)
    p = right.add_paragraph()
    p.add_run(f"Date: {c.report_date}").bold = True
    if c.is_sample:
        p = right.add_paragraph()
        r = p.add_run("SAMPLE DATA – NOT FOR RELEASE")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(8), RGBColor.from_string("A61B1B")
    _gap(doc, 6)

    info_w = [CONTENT_W_MM * f for f in (0.17, 0.33, 0.15, 0.35)]
    _section(doc, "Report Information")
    _label_value_table(doc, c.info_rows, info_w)
    _gap(doc, 6)

    details = []
    if c.tensile_details:
        details.append(f"Tensile: {c.tensile_details}")
    if c.hardness_details:
        details.append(f"Hardness: {c.hardness_details}")
    if c.criteria_source:
        details.append(f"Acceptance criteria source: {c.criteria_source}")
    _section(doc, "Mechanical Testing")
    _results_table(doc, c.mechanical, details)
    _gap(doc, 6)

    _section(doc, "Metallography")
    _results_table(doc, c.metallography, [])
    _label_value_table(doc, c.metallography_rows, info_w)
    wide = [("Preparation", c.preparation)] if c.preparation else []
    wide.append(("Microstructure Observation", c.microstructure))
    _label_value_table(doc, wide, [info_w[0], CONTENT_W_MM - info_w[0]])
    p = doc.add_paragraph()
    r = p.add_run(c.footnote)
    r.font.size, r.font.color.rgb = Pt(6.5), MUTED
    _gap(doc, 4)

    _section(doc, "Inspection Summary")
    t = _table(doc, 1, [CONTENT_W_MM])
    _write(t.cell(0, 0), c.summary)
    if c.remarks:
        _gap(doc, 3)
        t = _table(doc, 1, [CONTENT_W_MM])
        _write(t.cell(0, 0), f"Remarks: {c.remarks}")
    _gap(doc, 6)

    fg, bg = STATUS_RGB[c.overall.value]
    t = _table(doc, 1, [CONTENT_W_MM * f for f in (0.14, 0.30, 0.56)])
    _shade(t.cell(0, 0), NAVY)
    _write(t.cell(0, 0), "OVERALL RESULT", bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _shade(t.cell(0, 1), bg)
    _write(t.cell(0, 1), c.overall.value, bold=True, size=15, color=RGBColor.from_string(fg),
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _write(t.cell(0, 2), "\n".join(c.overall_reasons[:4]), size=7.5)
    _gap(doc, 6)

    sig_w = [CONTENT_W_MM * f for f in (0.12, 0.38, 0.12, 0.38)]
    t = _table(doc, 4, sig_w)
    a = t.cell(0, 0).merge(t.cell(0, 1))
    b = t.cell(0, 2).merge(t.cell(0, 3))
    for cell, text in ((a, "INSPECTED BY"), (b, "REVIEWED BY")):
        _shade(cell, LABEL_BG)
        _write(cell, text, bold=True, size=7.5, color=MUTED)
    for r, label in enumerate(("Name", "Signature", "Date"), start=1):
        for col in (0, 2):
            _shade(t.cell(r, col), LABEL_BG)
            _write(t.cell(r, col), label, bold=True, size=7.5, color=MUTED)
    _write(t.cell(1, 1), c.inspector)
    t.rows[2].height = Mm(9)

    footer = sec.footer.paragraphs[0]
    note = "Test status calculated deterministically from stated criteria"
    if c.ai_used:
        note += " · AI-assisted wording reviewed by inspector"
    r = footer.add_run(f"Form {settings.FORM_NUMBER}    ·    {note}")
    r.font.size, r.font.color.rgb = Pt(6.5), MUTED
    return doc


def generate_docx(report: InspectionReport, evaluation: EvaluationResult) -> bytes:
    buf = io.BytesIO()
    _build(build_content(report, evaluation)).save(buf)
    return buf.getvalue()
