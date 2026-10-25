"""Covering letters: wording, documents, packages, history and integrity."""

import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfReader

from database.database import LetterFinalizedError, ReportRepository
from models.letter import STATUS_DRAFT, STATUS_FINAL, CompanyProfile, CoveringLetter
from models.sample_data import sample_report
from reports.letter_content import build_letter_content, outcome_words
from services.letter_service import (
    build_letter_documents,
    build_package,
    finalize_submission,
    find_reports,
    load_company,
    format_letter_number,
    next_letter_number,
    parse_letter_number,
    save_company,
    validate_letter,
)


def _report(number="QC-1", heat="HT-1", sample="S-1", date="2026-10-20"):
    r = sample_report()
    r.info.is_sample = False
    r.info.report_number, r.info.heat_batch, r.info.sample_number = number, heat, sample
    r.info.report_date = date
    return r


def _letter(**kw):
    base = dict(letter_number="AHP/L/001", letter_date="2026-10-25",
                recipient_organization="Project Office", report_numbers=["QC-1"])
    base.update(kw)
    return CoveringLetter(**base)


def _v1_letter(**kw):
    """Letter on the original template (previously issued letters)."""
    return _letter(template_version=1, **kw)


def _pdf_text(data: bytes) -> tuple[int, str]:
    reader = PdfReader(io.BytesIO(data))
    return len(reader.pages), "\n".join(p.extract_text() for p in reader.pages)


# --- Wording: original template (v1), kept for previously issued letters --------

def test_v1_indenter_and_recipient_are_dynamic():
    for indenter, recipient in (("PCSIR", "Project Office A"), ("Client Ltd", "Directorate B")):
        c = build_letter_content(_v1_letter(indenter=indenter, recipient_organization=recipient),
                                 [_report()], CompanyProfile())
        body = " ".join(c.paragraphs)
        assert f"received from {indenter}" in body
        assert f"onward transmission to {indenter}" in body
        assert recipient in c.recipient_lines


def test_v1_forward_to_takes_precedence_for_onward_transmission():
    c = build_letter_content(_v1_letter(indenter="PCSIR", forward_to="Chief Engineer, Project XYZ"),
                             [_report()], CompanyProfile())
    assert "onward transmission to Chief Engineer, Project XYZ" in " ".join(c.paragraphs)


def test_v1_missing_details_are_omitted_not_invented():
    r = _report()
    r.info.customer = ""  # no indenter can be derived either
    c = build_letter_content(_v1_letter(), [r], CompanyProfile())
    body = " ".join(c.paragraphs)
    assert "vide reference" not in body
    assert "received from" not in body
    assert "onward transmission" not in body
    assert c.signatory_name == ""
    assert c.subject.startswith("Submission of Inspection and Testing Report")


def test_letter_never_states_outcomes():
    failing = _report()
    failing.tensile.uts = 1  # would FAIL
    c = build_letter_content(_letter(project_name="P"), [failing], CompanyProfile())
    assert outcome_words(" ".join(c.paragraphs) + c.subject) == []


def test_activities_come_from_report_data():
    r = _report()
    r.metallography.grain_size, r.metallography.microstructure = None, ""
    r.hardness.readings = []
    body = " ".join(build_letter_content(_letter(), [r], CompanyProfile()).paragraphs)
    assert "tensile testing" in body
    assert "hardness" not in body and "metallographic" not in body


def test_v1_auto_values_from_reports_and_manual_override():
    reports = [_report("QC-1", "HT-1", "S-1", "2026-10-20"), _report("QC-2", "HT-2", "S-2", "2026-10-22")]
    c = build_letter_content(_v1_letter(report_numbers=["QC-1", "QC-2"]), reports, CompanyProfile())
    body = " ".join(c.paragraphs)
    assert "Sample No. S-1, S-2; Heat/Batch HT-1, HT-2" in body
    assert "20 October 2026 to 22 October 2026" in body
    c2 = build_letter_content(_v1_letter(report_numbers=["QC-1", "QC-2"], testing_date="21–23 Oct 2026"),
                              reports, CompanyProfile())
    assert "21–23 Oct 2026" in " ".join(c2.paragraphs)


def test_v1_multiple_reports_listed_in_order():
    reports = [_report("QC-2", "HT-2"), _report("QC-1", "HT-1")]
    c = build_letter_content(_v1_letter(report_numbers=["QC-2", "QC-1"]), reports, CompanyProfile())
    assert "Inspection Reports Nos. QC-2 and QC-1" in " ".join(c.paragraphs)
    assert c.enclosures[0].startswith("Mechanical Testing and Metallography Inspection Report No. QC-2")
    assert c.enclosures[1].startswith("Mechanical Testing and Metallography Inspection Report No. QC-1")


def test_body_override_used_but_enclosures_stay_deterministic():
    c = build_letter_content(_letter(body_override="First paragraph.\n\nSecond paragraph."),
                             [_report()], CompanyProfile())
    assert c.paragraphs == ["First paragraph.", "Second paragraph."]
    assert "No. QC-1" in c.enclosures[0]


# --- Validation ----------------------------------------------------------------

def test_required_letter_fields():
    errors, _ = validate_letter(CoveringLetter(), [])
    assert "Letter Number / Reference Number is required." in errors
    assert "Letter Date is required." in errors
    assert "Recipient organization is required." in errors
    assert "Attach at least one inspection report." in errors


def test_missing_report_is_an_error(tmp_path):
    repo = ReportRepository(tmp_path / "h.db")
    reports, missing = find_reports(repo, ["NOPE"])
    errors, _ = validate_letter(_letter(report_numbers=["NOPE"]), reports, missing)
    assert any("not found in history: NOPE" in e for e in errors)


def test_outcome_words_in_edited_wording_warn():
    _, warnings = validate_letter(_letter(body_override="All samples passed."), [_report()])
    assert any("mentions results" in w for w in warnings)


# --- Documents ------------------------------------------------------------------

def test_letter_pdf_one_page_with_number_and_date():
    pdf, docx, pages = build_letter_documents(_letter(), [_report()], CompanyProfile())
    n, text = _pdf_text(pdf.data)
    assert n == pages == 1
    assert "AHP/L/001" in text
    assert "25 October 2026" in text
    assert "AH Precision Inspection & Testing Services" in text
    assert "QC-1" in text


def test_long_letter_still_one_page():
    reports = [_report(f"QC-{i}", f"HT-{i}", f"S-{i}") for i in range(1, 6)]
    letter = _letter(report_numbers=[r.info.report_number for r in reports],
                     recipient_address="Building 1\nStreet 2\nSector 3\nCity 4",
                     additional_references="Ref A: 123\nRef B: 456", project_name="Long Project Name",
                     project_reference="PRJ-2026-0001", indenter="PCSIR", submission_reference="S/45",
                     receipt_date="2026-10-18", sample_count=5, signatory_name="Name")
    _, _, pages = build_letter_documents(letter, reports, CompanyProfile(address="A\nB", email="e@x"))
    assert pages == 1


def test_letter_docx_contents():
    _, docx, _ = build_letter_documents(_letter(forward_to="PCSIR"), [_report()], CompanyProfile())
    text = "\n".join(p.text for p in Document(io.BytesIO(docx.data)).paragraphs)
    assert "AHP/L/001" in text and "25 October 2026" in text and "PCSIR" in text
    assert "Enclosures:" in text and "QC-1" in text


def test_package_contains_all_documents_and_zip():
    reports = [_report("QC-1"), _report("QC-2", "HT-2")]
    pkg = build_package(_letter(report_numbers=["QC-1", "QC-2"]), reports, CompanyProfile())
    names = [f.filename for f in pkg.all_files]
    assert names == [
        "Covering_Letter_AHP_L_001.pdf", "Covering_Letter_AHP_L_001.docx",
        "Inspection_Report_QC-1.pdf", "Inspection_Report_QC-1.docx",
        "Inspection_Report_QC-2.pdf", "Inspection_Report_QC-2.docx",
    ]
    with zipfile.ZipFile(io.BytesIO(pkg.zip_file.data)) as zf:
        assert sorted(zf.namelist()) == sorted(names)
    # the enclosed report PDFs are the right reports
    assert "QC-2" in _pdf_text(pkg.report_files[2].data)[1]


# --- History and integrity ---------------------------------------------------

def test_finalized_letter_is_preserved_and_not_overwritten(tmp_path):
    repo = ReportRepository(tmp_path / "h.db")
    report = _report()
    repo.save(report, "PASS")
    letter = _letter(project_name="Original")
    repo.save_letter(letter)
    assert repo.letter_status("AHP/L/001") == STATUS_DRAFT
    finalize_submission(repo, letter, [report], CompanyProfile())
    assert repo.letter_status("AHP/L/001") == STATUS_FINAL

    with pytest.raises(LetterFinalizedError):
        repo.save_letter(_letter(project_name="Changed"))
    with pytest.raises(LetterFinalizedError):
        finalize_submission(repo, _letter(project_name="Changed"), [report], CompanyProfile())

    # Later edits to the report do not alter the finalized snapshot.
    report.tensile.uts = 999
    repo.save(report, "PASS")
    record = repo.load_letter(repo.search_letters("AHP/L/001")[0]["id"])
    assert record["letter"].project_name == "Original"
    assert record["reports"][0].tensile.uts == 325
    assert record["company"].name == "AH Precision Inspection & Testing Services"

    repo.save_letter(_letter(letter_number="QC-AH-OL-05/26"))
    assert len(repo.search_letters()) == 2


def test_letter_search_and_previous_values(tmp_path):
    repo = ReportRepository(tmp_path / "h.db")
    repo.save_letter(_letter(letter_number="L-1", indenter="PCSIR", report_numbers=["QC-7"]))
    repo.save_letter(_letter(letter_number="L-2", indenter="Client Ltd", report_numbers=["QC-8"]))
    assert [r["letter_number"] for r in repo.search_letters("QC-7")] == ["L-1"]
    assert [r["letter_number"] for r in repo.search_letters("pcsir")] == ["L-1"]
    assert set(repo.letter_field_values("indenter")) == {"PCSIR", "Client Ltd"}


def test_company_profile_defaults_and_persistence(tmp_path):
    repo = ReportRepository(tmp_path / "h.db")
    profile = load_company(repo)
    assert profile.name == "AH Precision Inspection & Testing Services"
    assert profile.tagline == "Precision • Integrity • Quality"
    assert profile.address == profile.telephone == profile.email == ""  # nothing invented
    profile.name = "Renamed Co."
    save_company(repo, profile)
    assert load_company(ReportRepository(tmp_path / "h.db")).name == "Renamed Co."
    profile.name = " "
    with pytest.raises(ValueError):
        save_company(repo, profile)


def test_existing_report_history_untouched_by_letters(tmp_path):
    repo = ReportRepository(tmp_path / "h.db")
    report = _report()
    repo.save(report, "PASS")
    finalize_submission(repo, _letter(), [report], CompanyProfile())
    assert repo.get_report("QC-1") == report
    assert [r["report_number"] for r in repo.search()] == ["QC-1"]


# --- Simplified template (v2) for new letters --------------------------------------

def test_new_letters_use_simplified_template_and_old_ones_keep_v1():
    assert CoveringLetter().template_version == 2
    assert CoveringLetter.from_dict({"letter_number": "OLD"}).template_version == 1


def test_v2_wording_uses_report_data_without_project_fields():
    reports = [_report("QC-1", "HT-1", "S-1"), _report("QC-2", "HT-2", "S-2")]
    c = build_letter_content(_letter(report_numbers=["QC-1", "QC-2"]), reports, CompanyProfile())
    body = " ".join(c.paragraphs)
    assert c.subject == "Submission of Inspection and Testing Reports"
    assert c.reference_lines == []
    assert body.startswith("The two samples of Aluminium Alloy Example Alloy "
                           "(Sample No. S-1, S-2; Heat/Batch HT-1, HT-2) have been received")
    assert "Inspection and Testing Reports Nos. QC-1 and QC-2 are enclosed herewith" in body
    assert "transmission to the concerned authority." in body
    assert c.enclosures == [
        "Inspection and Testing Report No. QC-1 (Aluminium Alloy Example Alloy, Heat/Batch HT-1)",
        "Inspection and Testing Report No. QC-2 (Aluminium Alloy Example Alloy, Heat/Batch HT-2)",
    ]
    for word in ("project", "indenter", "vide reference", "received from"):
        assert word not in body.lower()


def test_v2_single_report_and_named_authority():
    c = build_letter_content(_letter(forward_to="PCSIR"), [_report()], CompanyProfile())
    body = " ".join(c.paragraphs)
    assert c.subject == "Submission of Inspection and Testing Report"
    assert "The sample of Aluminium Alloy Example Alloy" in body and "has been received" in body
    assert "Inspection and Testing Report No. QC-1 is enclosed herewith" in body
    assert "transmission to PCSIR." in body


def test_v2_manual_sample_count_and_ids_take_precedence():
    c = build_letter_content(_letter(sample_count=4, sample_ids="Bars B1–B4"), [_report()],
                             CompanyProfile())
    assert "The four samples of Aluminium Alloy Example Alloy (Bars B1–B4)" in c.paragraphs[0]


def test_v2_missing_material_is_not_invented():
    r = _report()
    r.info.material = r.info.grade = r.info.sample_number = r.info.heat_batch = ""
    c = build_letter_content(_letter(), [r], CompanyProfile())
    assert c.paragraphs[0].startswith(
        "The samples pertaining to the material/product identified in the enclosed inspection report "
        "have been received")


def test_v2_letter_fits_one_page_with_many_reports():
    reports = [_report(f"QC-{i}", f"HT-{i}", f"S-{i}") for i in range(1, 7)]
    letter = _letter(report_numbers=[r.info.report_number for r in reports],
                     recipient_name="The Director", recipient_department="Quality Division",
                     recipient_address="Building 1\nStreet 2\nCity", signatory_name="Name")
    _, _, pages = build_letter_documents(letter, reports, CompanyProfile(address="A", email="e@x"))
    assert pages == 1


# --- Letter numbering -----------------------------------------------------------

def test_letter_number_format_and_parsing():
    assert format_letter_number(1, 2026) == "QC-AH-OL-01/26"
    assert format_letter_number(12, 27) == "QC-AH-OL-12/27"
    assert parse_letter_number("qc-ah-ol-07/26") == (7, 26)
    assert parse_letter_number("QC-AH-OL-01") is None  # legacy number without year


def test_next_letter_number_avoids_duplicates(tmp_path):
    from datetime import date

    repo = ReportRepository(tmp_path / "h.db")
    on = date(2026, 10, 25)
    assert next_letter_number(repo, on) == "QC-AH-OL-01/26"
    repo.save_letter(_letter(letter_number="QC-AH-OL-01"))  # legacy, no year suffix (dated 2026)
    assert next_letter_number(repo, on) == "QC-AH-OL-02/26"
    repo.save_letter(_letter(letter_number="QC-AH-OL-02/26"))
    repo.save_letter(_letter(letter_number="QC-AH-OL-07/25"))  # other year
    assert next_letter_number(repo, on) == "QC-AH-OL-03/26"
    assert next_letter_number(repo, date(2027, 1, 2)) == "QC-AH-OL-01/27"
